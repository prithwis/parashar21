# 𝓟𝓪𝓻𝓪𝓼𝓱𝓪𝓻21
# 𝓟𝓻𝓲𝓽𝓱𝔀𝓲𝓼 𝓜𝓾𝓴𝓮𝓻𝓳𝓮𝓮
# --------------------------------------------------
# Global Variables
import p21
# --------------------------------------------------
#Utility functions 

import pandas as pd
import dateutil
import matplotlib.pyplot as plt
import math
import math
import numbers
#import collections
#import string_utils

import json

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from datetime import timedelta

#import swisseph as swe
import pytz



'''
Functions for Retrieval and Reporting



'''

#
# --------------------------------------------------
#
# Generates the text that appears in a chart image
#
def generateChartTxt():
    global txt
    txt = ['']*13
    for g,r in p21.GRashiN.items():
        if r == 1:
            txt[1]= txt[1]+' '+g+('/R' if p21.GRet[g] else '')
        if r == 2:
            txt[2]= txt[2]+' '+g+('/R' if p21.GRet[g] else '')
        if r == 3:
            txt[3]= txt[3]+' '+g+('/R' if p21.GRet[g] else '')
        if r == 4:
            txt[4]= txt[4]+' '+g+('/R' if p21.GRet[g] else '')
        if r == 5:
            txt[5]= txt[5]+' '+g+('/R' if p21.GRet[g] else '')
        if r == 6:
            txt[6]= txt[6]+' '+g+('/R' if p21.GRet[g] else '')
        if r == 7:
            txt[7]= txt[7]+' '+g+('/R' if p21.GRet[g] else '')
        if r == 8:
            txt[8]= txt[8]+' '+g+('/R' if p21.GRet[g] else '')
        if r == 9:
            txt[9]= txt[9]+' '+g+('/R' if p21.GRet[g] else '')
        if r == 10:
            txt[10]= txt[10]+' '+g+('/R' if p21.GRet[g] else '')
        if r == 11:
            txt[11]= txt[11]+' '+g+('/R' if p21.GRet[g] else '')
        if r == 12:
            txt[12]= txt[12]+' '+g+('/R' if p21.GRet[g] else '')
            
    for i in range(1,13):
        if len(txt[i]) == 0:
            txt[i] = '*'
#print(txt[1],txt[2],txt[3],txt[4],txt[5],txt[6],txt[7],txt[8],txt[9],txt[10],txt[11],txt[12],)

# --------------------------------------------------
# Plots the chart in the Bengal Format
#
def R12B_drawChart_Bengal():

    #ChartFile = p21.pName+p21.ChartType+'.png'
    #ChartFile = 'CurrentChart.png'
    generateChartTxt()

    #id = ChartType+'\n'+pName+'\n'+pDate+'\n'+pTime+'\n'+pPlace
    id = p21.ChartType+'\n'+p21.pName
    if p21.ChartType == 'Rashi':
        ChartColour = 'orange'
        ChartFile = 'RashiChart.png'
    else:
        ChartColour = 'olive'
        ChartFile = 'NavamsaChart.png'
    
        
    #plt.figure(figsize=(7,7))
    plt.figure(figsize=(7,7),facecolor=ChartColour)
    
    plt.axis('off')

    # draw vertical line 
    plt.plot([30, 30], [0, 90], 'k-', lw=2)
    plt.plot([60, 60], [0, 90], 'k-', lw=2)

    # draw horizontal line 
    plt.plot([0, 90], [30, 30], 'k-', lw=2)
    plt.plot([0, 90], [60, 60], 'k-', lw=2)

    #draw diagonal lines
    plt.plot([60,90],[60,90], 'k-', lw=2)
    plt.plot([0,30],[90,60], 'k-', lw=2)
    plt.plot([0,30],[90,60], 'k-', lw=2)
    plt.plot([0,30],[0,30], 'k-', lw=2)
    plt.plot([60,90],[30,0], 'k-', lw=2)

    plt.text(32, 38, id, fontsize=12)

    plt.text(32, 82, txt[1], fontsize=12)
    plt.text(8, 82, txt[2], fontsize=12)
    plt.text(2, 62, txt[3], fontsize=12)
    plt.text(2, 45, txt[4], fontsize=12)
    plt.text(2, 25, txt[5], fontsize=12)
    plt.text(8, 5, txt[6], fontsize=12)
    plt.text(32, 5, txt[7], fontsize=12)
    plt.text(62, 5, txt[8], fontsize=12)
    plt.text(68, 25, txt[9], fontsize=12)
    plt.text(68, 45, txt[10], fontsize=12)
    plt.text(68, 62, txt[11], fontsize=12)
    plt.text(62, 82, txt[12], fontsize=12)

    
    #plt.savefig("CurrentChart.png", bbox_inches='tight')
    plt.savefig(ChartFile, bbox_inches='tight')
    #plt.show()

def R01_CreateReportDoc(cqs):
      

    p21.document = Document()
    section = p21.document.sections[0]
    header = section.header
    footer = section.footer
    
    header01 = header.paragraphs[0]
    header01.text = "Parashar21 | Khona21 MongoDB database"

    now = datetime.now(pytz.timezone('Asia/Kolkata'))
    #print(now.strftime("%d %b %Y"))

    footer01 = footer.paragraphs[0]
    footer01.text = "Printed on : "+now.strftime("%d %b %Y")+"\nhttp://parashar21.blogspot.com | https://github.com/prithwis/parashar21"

    #heading_1 = p21.pName +" >>> "+p21.ChartType
    #heading_1 = ChartType+" Chart of "+pName
    
    heading_1 = "Selected Charts"  #p21.ReportFile
    p21.document.add_heading(heading_1, 0)
    heading_2 = cqs                         # current query string
    p21.document.add_heading(heading_2, level=3)
    
    p21.document.add_picture('./Saraswati.png', width=Inches(4.25))
    last_paragraph = p21.document.paragraphs[-1] 
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    threeLineGap = '\n\n\n'
    
    p21.document.add_paragraph(threeLineGap)
    
    gyan = "Astrology is the science of correlation, not causation."
    para = p21.document.add_paragraph(gyan)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p21.document.add_paragraph(threeLineGap)
    
    
    refer = "Astrology - An Application of Data Science \nhttps://bit.ly/pmastro"
    para = p21.document.add_paragraph(refer)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    
    #p21.document.add_page_break()
    #return document
    
# --------------------------------------------------

def R13C_ListPositions(text,L):
    # 
    L1 = L[1:]
    D = ''
    for c,e in enumerate(L1,1):
        if e :
            D = D+' '+str(c)+' ' 
    if len(D)>0:
        return(text+D)
    else:
        return(text+'<none>')    
    #return(D)

def R13B_ListPositions(text,L):
    # 
    L1 = L[1:]
    D = text+' : '
    for c,e in enumerate(L1,1):
        D = D+' '+str(c)+':'+str(e)+' '
    return(D)
        

def R13A_ShowTrueDict(desc,tDict):
    # creates a string with Keys that have Value True
    
    D = ''
    for (k,v) in tDict.items():
        if v :
            D = D+' '+k
    if len(D)>0:
        return(desc+D)
    else:
        return(desc+'<none>')
    #return(D)


def R511_parseChartData(c):

    p21.GLon = c['GLon']
    p21.GRet = c['GRet']
    pid = c['pid']
    p21.pName = pid['name']
    p21.GrahaLordBhav = c['GrahaLordBhav']
    p21.Lord = c['Lord']
    p21.GrahaBhava = c['GrahaBhava']
    p21.LordBhav = c['LordBhav']
    p21.exaltG = c['exaltG']
    p21.exaltL = c['exaltL']
    p21.debilG = c['debilG']
    p21.debilL = c['debilL']
    p21.inFriendG = c['inFriendG']
    p21.inFriendL = c['inFriendL']
    p21.inEnemyG = c['inEnemyG']
    p21.inEnemyL = c['inEnemyL']
    
    p21.GAspects2 = c['GAspects2']
    p21.GAspectedBy2 = c['GAspectedBy2']
    p21.BAspectedBy2 = c['BAspectedBy2']
    

def R512_FormatPage():
       
    p21.document.add_page_break()
    
    p2 = p21.document.add_paragraph()
    run_1 = p2.add_run()
    run_1.add_picture('./RashiChart.png', width=Inches(3.0))
    run_2 = p2.add_run()
    run_2.add_picture('./NavamsaChart.png', width=Inches(3.0))
    #p1 = p21.document.add_paragraph(p21.pName)                                             # skip printing name
    
    p21.document.add_paragraph(R13A_ShowTrueDict('Retrograde Graha',p21.GRet))     # Show Grahas that are retrograde

    cPara = R13B_ListPositions('Lord of ',p21.Lord)                                # Show Lords
    cPara = cPara+"\n"
    cPara = cPara+'Graha Lord of \n'+json.dumps(p21.GrahaLordBhav)                          # Show the Bhavs of whicha Graha is Lord
    p21.document.add_paragraph(cPara)

    p21.document.add_paragraph(R13B_ListPositions('Lord in Bhava',p21.LordBhav))   # Which Lord in which Bhav
            
    table = p21.document.add_table(rows=1, cols=2)
    cell1 = table.cell(0, 0)
    cell2 = table.cell(0, 1)
    
    
    cPara1 = 'Graha Status\n'
    cPara1 = cPara1+R13A_ShowTrueDict('Exalted     : ',p21.exaltG)+'\n'
    cPara1 = cPara1+R13A_ShowTrueDict('Debilitated : ',p21.debilG)+'\n'
    cPara1 = cPara1+R13A_ShowTrueDict('in Friend   : ',p21.inFriendG)+'\n'
    cPara1 = cPara1+R13A_ShowTrueDict('in Enemy    : ',p21.inEnemyG)
    cell1.add_paragraph(cPara1)
    
    
    cPara2 = 'Lord Status\n'
    cPara2 = cPara2+R13C_ListPositions('Exalted     : ',p21.exaltL)+'\n'
    cPara2 = cPara2+R13C_ListPositions('Debilitated : ',p21.debilL)+'\n'
    cPara2 = cPara2+R13C_ListPositions('In Friend   : ',p21.inFriendL)+'\n'
    cPara2 = cPara2+R13C_ListPositions('In Enemy    : ',p21.inEnemyL)
    cell2.add_paragraph(cPara2)
    #p21.document.add_paragraph(cPara2)
    
    
    cPara = 'Graha Aspects'
    cPara = cPara+'\nAspects🤓 '+json.dumps(p21.GAspects2)  
    cPara = cPara+'\nAspected BY 👀'+json.dumps(p21.GAspectedBy2)  
    p21.document.add_paragraph(cPara)
    
    #print(p21.GAspects2)
    #print(p21.GAspectedBy2)
    sorted_BAspectedBy2={}
    print(p21.BAspectedBy2)
    T1 = {}
    for k,v in p21.BAspectedBy2.items():
        T1[int(k)] =v
    T2 = dict(sorted(T1.items()))
    #print(T2)
    
    cPara = 'Bhav Aspect'
    cPara = cPara+'\nAspected BY 👀'+json.dumps(T2)  
    p21.document.add_paragraph(cPara)


    
    #p21.document.add_page_break()

def tracer():
    print('tracer')
    print(p21.ChartType)
    
print('p21utilsR imported')
